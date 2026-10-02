"""Build and validate Estime's approved local-listing URL registry.

Usage:
    python analytics/utm_registry.py refresh
    python analytics/utm_registry.py validate

The calendar workbook is the source for its 52 scheduled GBP post URLs.
Other channels are templates until Estime actually publishes a link.
"""

from __future__ import annotations

import json
import re
import sys
from datetime import date
from pathlib import Path
from urllib.parse import parse_qs, urlparse

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = ROOT / "seo" / "Calendario_Estimes_Cafe.xlsx"
REGISTRY = ROOT / "analytics" / "utm_registry.json"
PROFILE_URL = "https://www.estimescafe.com/?utm_source=google&utm_medium=organic&utm_campaign=gbp"
GBP_MENU_URL = "https://www.estimescafe.com/menu?utm_source=google&utm_medium=organic&utm_campaign=gbp&utm_content=menu"
APPLE_BUSINESS_CONNECT_URL = "https://www.estimescafe.com/?utm_source=apple_maps&utm_medium=organic&utm_campaign=apple_business_connect"
BING_PLACES_URL = "https://www.estimescafe.com/?utm_source=bing_places&utm_medium=organic&utm_campaign=bing_places"
BING_PLACES_MENU_URL = "https://www.estimescafe.com/menu?utm_source=bing_places&utm_medium=organic&utm_campaign=bing_places&utm_content=menu"
BING_PLACES_ORDER_ONLINE_URL = "https://www.estimescafe.com/order-online?utm_source=bing_places&utm_medium=organic&utm_campaign=bing_places&utm_content=order_online"
SOURCE_LEGACY = "google_business_profile"
SOURCE_CURRENT = "google"
GBP_LOCATION = "locations/16081983341986978598"
ALLOWED_KEYS = {"utm_source", "utm_medium", "utm_campaign", "utm_content", "utm_term"}
VALUE_RE = re.compile(r"^[a-z0-9]+(?:_[a-z0-9]+)*$")


def _read_calendar() -> list[dict]:
    wb = load_workbook(WORKBOOK, read_only=True, data_only=False)
    try:
        sheet = wb["Calendario 52 semanas"]
        rows = []
        for row in range(5, 57):
            number = row - 4
            when, title, button, url, status = (
                sheet.cell(row, 1).value,
                sheet.cell(row, 5).value,
                sheet.cell(row, 8).value,
                sheet.cell(row, 9).value,
                sheet.cell(row, 11).value,
            )
            rows.append({
                "week": number,
                "date": when.date().isoformat(),
                "title": title,
                "button": button,
                "url": url,
                "status": status,
            })
        return rows
    finally:
        wb.close()


def _record(record_id: str, channel: str, placement: str, status: str, created_by: str, url: str, date_value: str = "") -> dict:
    parsed = urlparse(url)
    utms = {key: values[0] for key, values in parse_qs(parsed.query).items() if key.startswith("utm_")}
    return {
        "id": record_id,
        "channel": channel,
        "placement": placement,
        "date": date_value,
        "status": status,
        "created_by": created_by,
        "destination": f"{parsed.scheme}://{parsed.netloc}{parsed.path}",
        "utms": utms,
        "url": url,
    }


def refresh() -> dict:
    other_channels = []
    previous_by_id: dict[str, dict] = {}
    if REGISTRY.exists():
        existing = json.loads(REGISTRY.read_text(encoding="utf-8"))
        previous_by_id = {item["id"]: item for item in existing.get("records", [])}
        managed_ids = {"gbp_profile_website", "gbp_profile_menu", "apple_business_connect_website", "bing_places_website", "bing_places_menu", "bing_places_order_online"}
        other_channels = [item for item in existing.get("records", []) if item.get("id") not in managed_ids and item.get("channel") != "gbp"]
    profile = _record(
        "gbp_profile_website", "gbp", "profile_website",
        "configured_pending_review", "unknown_preexisting", PROFILE_URL,
    )
    menu = _record(
        "gbp_profile_menu", "gbp", "profile_menu",
        "ready_to_configure", "Codex_updated_2026_10_01", GBP_MENU_URL,
    )
    apple = _record(
        "apple_business_connect_website", "apple_maps", "profile_website",
        "ready_to_configure", "Codex_updated_2026_10_01", APPLE_BUSINESS_CONNECT_URL,
    )
    bing = _record(
        "bing_places_website", "bing_places", "profile_website",
        "ready_to_configure", "Codex_updated_2026_10_01", BING_PLACES_URL,
    )
    bing_menu = _record(
        "bing_places_menu", "bing_places", "profile_menu",
        "ready_to_configure", "Codex_updated_2026_10_01", BING_PLACES_MENU_URL,
    )
    bing_order_online = _record(
        "bing_places_order_online", "bing_places", "profile_order_online",
        "ready_to_configure", "Codex_updated_2026_10_01", BING_PLACES_ORDER_ONLINE_URL,
    )
    for item in (profile, menu, apple, bing, bing_menu, bing_order_online):
        previous = previous_by_id.get(item["id"])
        if previous:
            item["status"] = previous["status"]
            item["created_by"] = previous["created_by"]
    records = [profile, menu]
    for row in _read_calendar():
        maps = urlparse(row["url"]).netloc == "www.google.com"
        record_id = f"gbp_week_{row['week']:02d}"
        record = _record(
            record_id,
            "gbp",
            "directions" if maps else "post_website",
            "published" if row["status"] == "Publicado" else "draft",
            "unknown_preexisting" if maps or row["status"] == "Publicado" else "Codex_updated_2026_10_01",
            row["url"], row["date"],
        )
        previous = previous_by_id.get(record_id)
        if previous:
            record["created_by"] = (
                previous["created_by"] if previous["url"] == row["url"]
                else "unknown_updated_from_calendar"
            )
        records.append(record)
    records.extend([apple, bing, bing_menu, bing_order_online, *other_channels])
    document = {
        "property": "properties/556917509",
        "gbp_location": GBP_LOCATION,
        "last_verified": date.today().isoformat(),
        "records": records,
    }
    validate(document)
    REGISTRY.write_text(json.dumps(document, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return document


def validate(document: dict | None = None) -> dict:
    if document is None:
        document = json.loads(REGISTRY.read_text(encoding="utf-8"))
    assert document["property"] == "properties/556917509"
    assert document["gbp_location"] == GBP_LOCATION
    records = document["records"]
    assert len(records) >= 58, f"Expected at least 58 records, got {len(records)}"
    by_id = {item["id"]: item for item in records}
    assert len(by_id) == len(records), "Duplicate registry IDs"
    assert by_id["gbp_profile_website"]["url"] == PROFILE_URL
    assert sum(item["channel"] == "gbp" for item in records) == 54
    seen_web_urls: set[str] = {PROFILE_URL, GBP_MENU_URL}
    counts = {"published_posts": 0, "draft_web_posts": 0, "directions": 0}
    for row in _read_calendar():
        item = by_id[f"gbp_week_{row['week']:02d}"]
        assert item["url"] == row["url"], f"Calendar/registry URL mismatch in week {row['week']}"
        assert item["date"] == row["date"]
        assert item["status"] == ("published" if row["status"] == "Publicado" else "draft")
        parsed = urlparse(item["url"])
        pairs = parse_qs(parsed.query, keep_blank_values=True)
        assert item["destination"] == f"{parsed.scheme}://{parsed.netloc}{parsed.path}"
        assert parsed.scheme == "https"
        if item["placement"] == "directions":
            counts["directions"] += 1
            assert parsed.netloc == "www.google.com"
            assert not item["utms"] and not any(key.startswith("utm_") for key in pairs)
            continue
        assert item["placement"] == "post_website"
        assert parsed.netloc == "www.estimescafe.com"
        assert item["url"] not in seen_web_urls, "Duplicate website URL"
        seen_web_urls.add(item["url"])
        assert all(len(values) == 1 for values in pairs.values()), "Duplicate URL parameter"
        assert set(item["utms"]) == {"utm_source", "utm_medium", "utm_campaign", "utm_content"}
        assert item["utms"] == {key: values[0] for key, values in pairs.items() if key.startswith("utm_")}
        assert item["utms"]["utm_source"] == (SOURCE_LEGACY if row["week"] <= 4 else SOURCE_CURRENT)
        assert item["utms"]["utm_medium"] == "organic"
        assert item["utms"]["utm_campaign"] == "gbp_post"
        assert item["utms"]["utm_content"].startswith(f"week_{row['week']:02d}_")
        assert all(key in ALLOWED_KEYS and VALUE_RE.fullmatch(value) for key, value in item["utms"].items())
        assert not any("@" in value for value in item["utms"].values()), "Email in UTM"
        if row["status"] == "Publicado":
            counts["published_posts"] += 1
        else:
            counts["draft_web_posts"] += 1
    main = by_id["gbp_profile_website"]
    assert main["destination"] == "https://www.estimescafe.com/"
    assert main["utms"] == {"utm_source": "google", "utm_medium": "organic", "utm_campaign": "gbp"}
    menu = by_id["gbp_profile_menu"]
    assert menu["placement"] == "profile_menu"
    assert menu["destination"] == "https://www.estimescafe.com/menu"
    assert menu["utms"] == {"utm_source": "google", "utm_medium": "organic", "utm_campaign": "gbp", "utm_content": "menu"}
    apple = by_id["apple_business_connect_website"]
    assert apple["channel"] == "apple_maps"
    assert apple["placement"] == "profile_website"
    assert apple["destination"] == "https://www.estimescafe.com/"
    assert apple["utms"] == {"utm_source": "apple_maps", "utm_medium": "organic", "utm_campaign": "apple_business_connect"}
    bing = by_id["bing_places_website"]
    assert bing["channel"] == "bing_places"
    assert bing["placement"] == "profile_website"
    assert bing["destination"] == "https://www.estimescafe.com/"
    assert bing["utms"] == {"utm_source": "bing_places", "utm_medium": "organic", "utm_campaign": "bing_places"}
    bing_menu = by_id["bing_places_menu"]
    assert bing_menu["channel"] == "bing_places"
    assert bing_menu["placement"] == "profile_menu"
    assert bing_menu["destination"] == "https://www.estimescafe.com/menu"
    assert bing_menu["utms"] == {"utm_source": "bing_places", "utm_medium": "organic", "utm_campaign": "bing_places", "utm_content": "menu"}
    bing_order_online = by_id["bing_places_order_online"]
    assert bing_order_online["channel"] == "bing_places"
    assert bing_order_online["placement"] == "profile_order_online"
    assert bing_order_online["destination"] == "https://www.estimescafe.com/order-online"
    assert bing_order_online["utms"] == {"utm_source": "bing_places", "utm_medium": "organic", "utm_campaign": "bing_places", "utm_content": "order_online"}
    for item in records:
        if item["channel"] == "gbp":
            continue
        parsed = urlparse(item["url"])
        pairs = parse_qs(parsed.query, keep_blank_values=True)
        assert parsed.scheme == "https" and parsed.netloc == "www.estimescafe.com"
        assert item["destination"] == f"{parsed.scheme}://{parsed.netloc}{parsed.path}"
        assert item["url"] not in seen_web_urls, "Duplicate website URL"
        seen_web_urls.add(item["url"])
        assert all(len(values) == 1 for values in pairs.values()), "Duplicate URL parameter"
        assert {"utm_source", "utm_medium", "utm_campaign"} <= set(item["utms"])
        assert item["utms"] == {key: values[0] for key, values in pairs.items() if key.startswith("utm_")}
        assert all(key in ALLOWED_KEYS and VALUE_RE.fullmatch(value) for key, value in item["utms"].items())
        assert item["created_by"] and item["placement"] and item["status"]
    assert counts == {"published_posts": 4, "draft_web_posts": 45, "directions": 3}, counts
    return {"records": len(records), **counts, "result": "valid"}


if __name__ == "__main__":
    operation = sys.argv[1] if len(sys.argv) > 1 else "validate"
    if operation == "refresh":
        refresh()
    elif operation != "validate":
        raise SystemExit("Use refresh or validate")
    print(json.dumps(validate(), ensure_ascii=False))
